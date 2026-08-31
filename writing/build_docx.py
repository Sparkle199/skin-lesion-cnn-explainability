"""Converts a writing/*.md chapter into a formatted .docx: 11pt body font,
1.5 line spacing, proper heading levels, and real Word tables for any
markdown table. Minimal, hand-rolled markdown handling scoped to what
these chapters actually use (#/##/###, **bold**, *italic*, pipe tables)
rather than a general-purpose parser.

Usage: python build_docx.py [chapter-stem]
  chapter-stem defaults to "chapter-3-methodology" if omitted.
"""

import base64
import hashlib
import re
import sys
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

MERMAID_CACHE = Path(__file__).parent / ".mermaid_cache"
MERMAID_CACHE.mkdir(exist_ok=True)


def render_mermaid(source: str) -> Path:
    """Renders mermaid diagram source to a cached PNG via mermaid.ink.
    Word cannot render mermaid natively, so the docx build needs an actual
    image; the .md source stays live-rendered mermaid for GitHub."""
    digest = hashlib.sha256(source.encode("utf-8")).hexdigest()[:16]
    cached = MERMAID_CACHE / f"{digest}.png"
    if cached.exists():
        return cached
    encoded = base64.urlsafe_b64encode(source.encode("utf-8")).decode("ascii")
    url = f"https://mermaid.ink/img/{encoded}?type=png&bgColor=white&width=1400"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        cached.write_bytes(resp.read())
    return cached

STEM = sys.argv[1] if len(sys.argv) > 1 else "chapter-3-methodology"
SRC = Path(__file__).parent / f"{STEM}.md"
OUT = Path(__file__).parent / f"{STEM}.docx"

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(11)


def set_default_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)

    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(10)

    for heading_name, size, bold in [
        ("Heading 1", 18, True),
        ("Heading 2", 15, True),
        ("Heading 3", 13, True),
    ]:
        style = doc.styles[heading_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    doc.sections[0].left_margin = Cm(2.5)
    doc.sections[0].right_margin = Cm(2.5)
    doc.sections[0].top_margin = Cm(2.5)
    doc.sections[0].bottom_margin = Cm(2.5)


INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")


def add_runs_with_inline_formatting(paragraph, text):
    """Splits text on **bold** and *italic* markers and adds runs accordingly."""
    tokens = []
    pos = 0
    pattern = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")
    for m in pattern.finditer(text):
        if m.start() > pos:
            tokens.append(("plain", text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(("bold", m.group(1)))
        else:
            tokens.append(("italic", m.group(2)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("plain", text[pos:]))

    for kind, chunk in tokens:
        run = paragraph.add_run(chunk)
        run.font.name = FONT_NAME
        run.font.size = FONT_SIZE
        if kind == "bold":
            run.bold = True
        elif kind == "italic":
            run.italic = True


IMAGE_LINE = re.compile(r"^!\[(.*?)\]\((.*?)\)$")


def add_figure(doc, caption, rel_path):
    add_figure_from_path(doc, caption, SRC.parent / rel_path)


def add_figure_from_path(doc, caption, img_path):
    max_width = Cm(14.5)
    max_height = Cm(21)
    try:
        from PIL import Image
        with Image.open(img_path) as im:
            px_w, px_h = im.size
        aspect = px_h / px_w
        width = max_width
        height = round(width * aspect)
        if height > max_height:
            height = max_height
            width = round(height / aspect)
    except Exception:
        width, height = max_width, None

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    if height is not None:
        run.add_picture(str(img_path), width=width, height=height)
    else:
        run.add_picture(str(img_path), width=width)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cap.paragraph_format.space_after = Pt(12)
    cap_run = cap.add_run(caption)
    cap_run.font.name = FONT_NAME
    cap_run.font.size = Pt(10)
    cap_run.italic = True


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table(lines, start_idx):
    rows = []
    i = start_idx
    while i < len(lines) and is_table_row(lines[i]):
        rows.append(lines[i])
        i += 1
    cells = [
        [c.strip() for c in row.strip().strip("|").split("|")]
        for row in rows
    ]
    if len(cells) >= 2 and re.match(r"^:?-+:?$", cells[1][0]):
        del cells[1]
    return cells, i


def add_table(doc, cells):
    n_rows = len(cells)
    n_cols = len(cells[0])
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = p.add_run(val)
            run.font.name = FONT_NAME
            run.font.size = Pt(10.5)
            if r == 0:
                run.bold = True
    doc.add_paragraph()


def build():
    text = SRC.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    set_default_style(doc)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(level=1)
            add_runs_with_inline_formatting(p, stripped[2:])
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_heading(level=2)
            add_runs_with_inline_formatting(p, stripped[3:])
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_heading(level=3)
            add_runs_with_inline_formatting(p, stripped[4:])
            i += 1
            continue

        if stripped.startswith("**Table"):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_runs_with_inline_formatting(p, stripped)
            i += 1
            continue

        if is_table_row(stripped):
            cells, i = parse_table(lines, i)
            add_table(doc, cells)
            continue

        img_match = IMAGE_LINE.match(stripped)
        if img_match:
            add_figure(doc, img_match.group(1), img_match.group(2))
            i += 1
            continue

        if stripped == "```mermaid":
            j = i + 1
            body = []
            while j < n and lines[j].strip() != "```":
                body.append(lines[j])
                j += 1
            source = "\n".join(body)
            i = j + 1

            caption = None
            if i < n and lines[i].strip() == "":
                i += 1
            if i < n:
                nxt = lines[i].strip()
                if nxt.startswith("*") and nxt.endswith("*") and not nxt.startswith("**"):
                    caption = nxt.strip("*")
                    i += 1

            print(f"Rendering mermaid diagram ({len(source)} chars)...")
            png_path = render_mermaid(source)
            add_figure_from_path(doc, caption or "", png_path)
            continue

        if stripped == "---":
            i += 1
            continue

        is_full_line_note = stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**")

        p = doc.add_paragraph()
        if STEM == "references" and is_full_line_note:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(stripped.strip("*"))
            run.font.name = FONT_NAME
            run.font.size = Pt(10)
            run.italic = True
        elif STEM == "references":
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-1.25)
            p.paragraph_format.space_after = Pt(10)
            add_runs_with_inline_formatting(p, stripped)
        else:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            add_runs_with_inline_formatting(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
